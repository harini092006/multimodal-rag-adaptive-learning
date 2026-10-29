"""Multimodal document ingestion.

Text PDF/DOCX/PPTX/TXT/MD  -> text, headings, tables
Scanned PDF / images        -> OCR (Tesseract via pytesseract)
Audio / video lectures      -> speech transcript (faster-whisper)

Every extractor returns structured segments:  {text, page_ref, section, modality}.
Missing optional engines produce a clear, user-readable error instead of silently returning nothing.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List

from app.config import settings

TEXT_EXT = {".txt", ".md", ".markdown"}
DOC_EXT = {".pdf", ".docx", ".pptx"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
MEDIA_EXT = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".mp4", ".mkv", ".webm", ".mov"}
ALLOWED_EXT = TEXT_EXT | DOC_EXT | IMAGE_EXT | MEDIA_EXT


class ParseError(Exception):
    pass


@dataclass
class ParseResult:
    segments: List[Dict[str, Any]] = field(default_factory=list)
    page_count: int = 0
    modalities: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(s["text"] for s in self.segments)

    @property
    def char_count(self) -> int:
        return sum(len(s["text"]) for s in self.segments)


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ---------------------------------------------------------------- OCR
def ocr_available() -> bool:
    if not settings.OCR_ENABLED:
        return False
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def ocr_languages() -> dict:
    """Requested (OCR_LANGS) vs installed Tesseract language packs. Only installed packs are ever used."""
    requested = [l for l in settings.OCR_LANGS.split("+") if l]
    try:
        import pytesseract
        installed = sorted(set(pytesseract.get_languages(config="")) - {"osd"})
    except Exception:
        installed = []
    usable = [l for l in requested if l in installed]
    return {"requested": requested, "installed": installed, "usable": usable,
            "missing": [l for l in requested if l not in installed]}


def ocr_image(img) -> str:
    try:
        import pytesseract
        langs = ocr_languages()["usable"]
        if not langs:
            raise RuntimeError("no requested language pack is installed")
        return clean_text(pytesseract.image_to_string(img, lang="+".join(langs)))
    except Exception as exc:
        raise ParseError(
            "OCR is not available on this server (install Tesseract and the 'pytesseract' package; "
            f"language packs for: {settings.OCR_LANGS}).") from exc


# ---------------------------------------------------------------- extractors
def _parse_text(path: Path, res: ParseResult) -> None:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    section = None
    buf: List[str] = []

    def flush():
        if buf:
            res.segments.append({"text": clean_text("\n".join(buf)), "section": section, "modality": "text"})
            buf.clear()

    for line in raw.splitlines():
        m = re.match(r"^\s{0,3}#{1,6}\s+(.*)$", line)
        if m and path.suffix.lower() in {".md", ".markdown"}:
            flush()
            section = m.group(1).strip()
            buf.append(section)
        else:
            buf.append(line)
    flush()
    res.page_count = 1
    res.modalities.append("text")


def _table_to_text(rows: List[List[str]]) -> str:
    lines = [" | ".join((c or "").strip().replace("\n", " ") for c in r) for r in rows if any((c or "").strip() for c in r)]
    return "Table:\n" + "\n".join(lines) if lines else ""


def _parse_pdf(path: Path, res: ParseResult) -> None:
    import pypdf
    try:
        reader = pypdf.PdfReader(str(path))
        if reader.is_encrypted:
            raise ParseError("The PDF is password protected.")
        pages = list(reader.pages)
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"Could not read the PDF ({type(exc).__name__}). The file may be corrupt.") from exc
    res.page_count = len(pages)

    tables_by_page: Dict[int, List[str]] = {}
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pl:
            for i, pg in enumerate(pl.pages[:200]):
                for t in pg.extract_tables() or []:
                    txt = _table_to_text(t)
                    if txt:
                        tables_by_page.setdefault(i, []).append(txt)
    except Exception:
        pass  # table extraction is best-effort

    pdfium_doc = None
    for i, page in enumerate(pages):
        try:
            text = clean_text(page.extract_text() or "")
        except Exception:
            text = ""
        if len(text) < 25:                          # likely a scanned page -> OCR
            if ocr_available():
                try:
                    import pypdfium2 as pdfium
                    pdfium_doc = pdfium_doc or pdfium.PdfDocument(str(path))
                    img = pdfium_doc[i].render(scale=2.5).to_pil()
                    otext = ocr_image(img)
                    if otext:
                        res.segments.append({"text": otext, "page_ref": f"p{i+1}", "modality": "ocr"})
                        if "ocr" not in res.modalities:
                            res.modalities.append("ocr")
                        continue
                except ParseError as exc:
                    res.warnings.append(str(exc))
                except Exception as exc:
                    res.warnings.append(f"OCR failed on page {i+1}: {type(exc).__name__}")
            else:
                res.warnings.append(f"Page {i+1} has no selectable text and OCR is unavailable.")
        if text:
            res.segments.append({"text": text, "page_ref": f"p{i+1}", "modality": "text"})
            if "text" not in res.modalities:
                res.modalities.append("text")
        for t in tables_by_page.get(i, []):
            res.segments.append({"text": t, "page_ref": f"p{i+1}", "modality": "table"})
            if "table" not in res.modalities:
                res.modalities.append("table")


def _parse_docx(path: Path, res: ParseResult) -> None:
    import docx
    try:
        d = docx.Document(str(path))
    except Exception as exc:
        raise ParseError(f"Could not read the DOCX ({type(exc).__name__}).") from exc
    section = None
    buf: List[str] = []

    def flush():
        if buf:
            res.segments.append({"text": clean_text("\n".join(buf)), "section": section, "modality": "text"})
            buf.clear()

    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if (p.style.name or "").lower().startswith("heading") or (p.style.name or "").lower() == "title":
            flush()
            section = t
        buf.append(t)
    flush()
    for tbl in d.tables:
        rows = [[c.text for c in r.cells] for r in tbl.rows]
        txt = _table_to_text(rows)
        if txt:
            res.segments.append({"text": txt, "modality": "table", "section": section})
            res.modalities.append("table")
    res.page_count = max(1, len(d.paragraphs) // 30)
    res.modalities.insert(0, "text")


def _parse_pptx(path: Path, res: ParseResult) -> None:
    from pptx import Presentation
    try:
        prs = Presentation(str(path))
    except Exception as exc:
        raise ParseError(f"Could not read the PPTX ({type(exc).__name__}).") from exc
    res.page_count = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        parts: List[str] = []
        title = None
        try:
            title = slide.shapes.title.text.strip() if slide.shapes.title is not None else None
        except Exception:
            title = None
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                t = _table_to_text(rows)
                if t:
                    res.segments.append({"text": t, "page_ref": f"slide {i}", "modality": "table", "section": title})
                    if "table" not in res.modalities:
                        res.modalities.append("table")
            if shape.shape_type == 13 and ocr_available():       # PICTURE
                try:
                    from PIL import Image
                    import io
                    otext = ocr_image(Image.open(io.BytesIO(shape.image.blob)))
                    if otext:
                        parts.append(otext)
                        if "ocr" not in res.modalities:
                            res.modalities.append("ocr")
                except Exception:
                    pass
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip())
        if parts:
            res.segments.append({"text": clean_text("\n".join(parts)), "page_ref": f"slide {i}", "section": title,
                                 "modality": "text"})
    if res.segments and "text" not in res.modalities:
        res.modalities.insert(0, "text")


def _parse_image(path: Path, res: ParseResult) -> None:
    if not ocr_available():
        raise ParseError("This is an image and OCR is not available on the server "
                         "(install Tesseract + pytesseract), so no text could be extracted.")
    from PIL import Image
    try:
        img = Image.open(path)
        img.load()
    except Exception as exc:
        raise ParseError("The image could not be opened.") from exc
    text = ocr_image(img)
    if not text:
        res.warnings.append("No readable text was found in the image.")
    else:
        res.segments.append({"text": text, "page_ref": "image", "modality": "ocr"})
    res.page_count = 1
    res.modalities.append("ocr")


def _parse_media(path: Path, res: ParseResult, language=None) -> None:
    from app.speech import stt
    try:
        tr = stt.transcribe_file(str(path), language)
    except stt.SpeechUnavailable as exc:
        raise ParseError(str(exc)) from exc
    if not tr.text:
        res.warnings.append("No speech was detected in the recording.")
    # group transcript into ~60-second windows so citations can point at a timestamp
    window, start, acc = 60.0, 0.0, []
    for seg in tr.segments:
        acc.append(seg["text"])
        if seg["end"] - start >= window:
            res.segments.append({"text": " ".join(acc), "page_ref": _fmt_ts(start), "modality": "transcript"})
            acc, start = [], seg["end"]
    if acc:
        res.segments.append({"text": " ".join(acc), "page_ref": _fmt_ts(start), "modality": "transcript"})
    res.page_count = 1
    res.modalities.append("transcript")


def _fmt_ts(sec: float) -> str:
    sec = int(sec)
    return f"{sec // 60:02d}:{sec % 60:02d}"


def parse_file(path: str, ext: str, language=None) -> ParseResult:
    p = Path(path)
    ext = ext.lower() if ext.startswith(".") else "." + ext.lower()
    if not p.exists():
        raise ParseError("The stored file could not be found.")
    res = ParseResult()
    if ext in TEXT_EXT:
        _parse_text(p, res)
    elif ext == ".pdf":
        _parse_pdf(p, res)
    elif ext == ".docx":
        _parse_docx(p, res)
    elif ext == ".pptx":
        _parse_pptx(p, res)
    elif ext in IMAGE_EXT:
        _parse_image(p, res)
    elif ext in MEDIA_EXT:
        _parse_media(p, res, language)
    else:
        raise ParseError(f"Unsupported file type '{ext}'.")
    if "ocr" in res.modalities:
        miss = ocr_languages()["missing"]
        if miss:
            res.warnings.append(f"OCR language pack(s) not installed: {', '.join(miss)}. Text in those languages was NOT read; "
                                "install the Tesseract language data and reprocess the material.")
    res.segments = [s for s in res.segments if s["text"].strip()]
    if not res.segments:
        detail = " ".join(res.warnings) or "No extractable text was found."
        raise ParseError(detail)
    return res
